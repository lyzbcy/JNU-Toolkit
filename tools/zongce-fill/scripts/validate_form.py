# -*- coding: utf-8 -*-
"""
校验已填写的《本科生综测个人加分项信息统计表》。
独立于 fill_form.py: 直接读 xlsx, 也可用于班长抽查同学手工填的表。
用法: python validate_form.py --file 已填表格.xlsx
退出码: 0=通过(可有WARN) / 2=有ERROR
"""
import argparse
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

COMP_CATEGORIES = ["算法与程序设计", "数据挖掘与人工智能", "信息可视化", "物联网与嵌入式系统",
                   "数学建模", "网络安全", "计算机体系结构", "项目应用与创新创业类"]
CERT_NAMES = ["CET-4", "CET-6", "雅思", "托福", "CSP", "计算机软考"]
PAPER_LEVELS = ["CCF-A类会议论文", "CCF-B类会议论文", "SCI/SCIE一区期刊论文", "SCI/SCIE二区期刊论文",
                "SCI/SCIE三区期刊论文", "SCI/SCIE四区期刊论文", "EI期刊论文", "SCI/SCIE收录的会议论文",
                "CSCD核心库期刊论文", "国家公布的核心期刊", "国际学术会议论文", "国内普通期刊",
                "国内学术会议发表的论文", "国家有关出版社公开出版的著作", "授权发明专利", "授理发明专利"]
DACHUANG_LEVELS = ["国家级大创创新项目", "国家级大创创业项目", "省级大创创新项目", "省级大创创业项目",
                   "校级大创", "院级大创"]


class Report:
    def __init__(self):
        self.errors, self.warns = [], []

    def error(self, where, msg):
        self.errors.append(f"[{where}] {msg}")

    def warn(self, where, msg):
        self.warns.append(f"[{where}] {msg}")


def row_vals(ws, r, ncols):
    return [ws.cell(row=r, column=c).value for c in range(1, ncols + 1)]


def nonempty_rows(ws, start, end, ncols):
    out = []
    for r in range(start, end + 1):
        vals = row_vals(ws, r, ncols)
        if any(v is not None and str(v).strip() != "" for v in vals):
            out.append((r, vals))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    args = ap.parse_args()
    wb = load_workbook(args.file)
    rep = Report()

    # ---- 示例红字检查: 模板示例行(例/XXX)是红字, 真实数据若继承红色格式则提示
    for wsx in wb.worksheets:
        if wsx.title == "Sheet1":
            continue
        for row in wsx.iter_rows(min_row=2):
            for cell in row:
                if cell.value is None:
                    continue
                fc = cell.font.color if cell.font else None
                if fc is not None and getattr(fc, "type", None) == "rgb" \
                        and str(fc.rgb).upper().endswith("FF0000"):
                    rep.warn("格式", f"{wsx.title}!{cell.coordinate} 数据为红色示例格式, 应为正常黑字")

    # 隐藏竞赛目录
    catalog = {}
    ws1 = wb["Sheet1"]
    for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, max_col=2):
        if row[0].value and row[1].value:
            catalog[str(row[0].value)] = str(row[1].value)

    # ---- 基础分
    ws = wb["基础分"]
    rows = nonempty_rows(ws, 4, ws.max_row, 6)
    if not rows:
        rep.error("基础分", "整表未填: 学号/姓名/学时/加分类型必填")
    else:
        sid, name, hours, ktype = rows[0][1][:4]
        if not sid or not name:
            rep.error("基础分", "学号或姓名为空")
        if not isinstance(hours, (int, float)):
            rep.error("基础分", "第二课堂学时必须是数字")
        elif hours > 40:
            rep.error("基础分", f"学时{hours}超过上限40")
        if ktype not in ("已参与两次讲座类", "已参与一次以上学科竞赛"):
            rep.error("基础分", f"加分类型[{ktype}]应为下拉二选一")
        elif ktype == "已参与一次以上学科竞赛":
            comp_cells = [(r, v) for r, v in rows if v[4]]
            if not comp_cells:
                rep.error("基础分", "选了学科竞赛但没填竞赛名称")
            for r, v in comp_cells:
                if str(v[4]) not in catalog:
                    rep.error("基础分", f"行{r} 竞赛[{v[4]}]不在目录中")

    # ---- 加分项-竞赛
    ws = wb["加分项-竞赛"]
    for r, v in nonempty_rows(ws, 3, ws.max_row, 9):
        name, level, alevel, agrade, rank, advisor, desc, cat, score = v
        if name not in catalog:
            rep.warn("竞赛", f"行{r} [{name}]不在学校竞赛目录, 待班长认定")
        else:
            if level != catalog[name]:
                rep.error("竞赛", f"行{r} 竞赛级别[{level}]与目录映射[{catalog[name]}]不符")
        if alevel not in ("国家级", "省级", "市级", "校级"):
            rep.error("竞赛", f"行{r} 获奖级别[{alevel}]非法")
        if agrade not in ("特等奖", "一等奖", "二等奖", "三等奖"):
            rep.error("竞赛", f"行{r} 获奖等级[{agrade}]非法")
        if cat not in COMP_CATEGORIES:
            rep.error("竞赛", f"行{r} 竞赛分类[{cat}]非法")
        if not desc or len(str(desc).strip()) < 10:
            rep.error("竞赛", f"行{r} 竞赛描述必填(写清赛道方向与个人工作)")
        if not isinstance(score, (int, float)):
            if agrade == "三等奖":
                rep.warn("竞赛", f"行{r} 三等奖档位未定, 分数留空待班长核定")
            elif level:
                rep.error("竞赛", f"行{r} 应得分数为空或非数字")
            else:
                rep.warn("竞赛", f"行{r} 待认定条目无分数, 正常")
        if agrade in ("特等奖", "三等奖"):
            rep.warn("竞赛", f"行{r} {agrade}档位口径需班长核定")

    # ---- 论文
    ws = wb["加分项-论文著作专利"]
    for r, v in nonempty_rows(ws, 3, ws.max_row, 7):
        typ, title, venue, level, rank, advisor, score = v
        if typ not in ("论文/著作", "专利"):
            rep.error("论文", f"行{r} 类别[{typ}]非法")
        if level not in PAPER_LEVELS:
            rep.error("论文", f"行{r} 级别[{level}]非法或使用了冗余值")
        if typ == "论文/著作" and not venue:
            rep.error("论文", f"行{r} 论文类必须填刊物/会议名称")
        if not isinstance(score, (int, float)):
            rep.error("论文", f"行{r} 应得分数为空或非数字")
        rep.warn("论文", f"行{r} 论文/收录认定需班长核定")

    # ---- 社会实践(双区)
    ws = wb["加分项-社会实践"]
    for r, v in nonempty_rows(ws, 3, 31, 5):
        name, tclass, role, desc, score = v
        if tclass and tclass not in ("省级以上重点团队", "校级重点团队", "院级重点团队", "院级一般团队"):
            rep.error("社会实践", f"行{r} 项目类别[{tclass}]非法")
        if role not in ("队长", "队员", "个人参与社会实践"):
            rep.error("社会实践", f"行{r} 团队任职[{role}]非法")
        if not desc or len(str(desc).strip()) < 20:
            rep.error("社会实践", f"行{r} 描述必填(50-100字)")
        if not isinstance(score, (int, float)):
            rep.error("社会实践", f"行{r} 应得分数为空或非数字")
    for r, v in nonempty_rows(ws, 32, ws.max_row, 5):
        level, b, role, title, score = v
        if level and level not in DACHUANG_LEVELS:
            rep.error("大创", f"行{r} 级别[{level}]非法")
        if b:
            rep.error("大创", f"行{r} B列应留空(大创不填团队类别)")
        if role not in ("队长", "队员", None):
            rep.error("大创", f"行{r} 任职[{role}]非法")
        if "未结题" in str(title or "") and score:
            rep.error("大创", f"行{r} 未结题不应有加分")

    # ---- 学生事务
    ws = wb["加分项-学生事务"]
    for r, v in nonempty_rows(ws, 3, ws.max_row, 4):
        pos, period, desc, score = v
        if not pos or not period:
            rep.error("学生事务", f"行{r} 职务/时间段必填")
        if not desc or len(str(desc).strip()) < 15:
            rep.error("学生事务", f"行{r} 描述必填(约50字)")
        if not isinstance(score, (int, float)) or not (0 <= score <= 5):
            rep.error("学生事务", f"行{r} 分数应在0-5")
        rep.warn("学生事务", f"行{r} 事务分以辅导员评定为准")

    # ---- 文体志愿
    ws = wb["加分项-文体志愿类"]
    for r, v in nonempty_rows(ws, 4, ws.max_row, 7):
        typ, name, alevel, agrade, hours, time, score = v
        if typ not in ("文艺类", "体育类", "志愿类", "其他荣誉"):
            rep.error("文体志愿", f"行{r} 类别[{typ}]非法")
        if typ in ("文艺类", "体育类", "其他荣誉") and alevel not in ("省部级以上", "市级", "校级", "院级"):
            rep.error("文体志愿", f"行{r} 获奖级别[{alevel}]非法")
        if typ == "志愿类" and not hours and "献血" not in str(name or ""):
            rep.error("文体志愿", f"行{r} 志愿类必须填服务时长")
        if not isinstance(score, (int, float)):
            rep.error("文体志愿", f"行{r} 应得分数为空或非数字")

    # ---- 其他
    ws = wb["加分项-其他"]
    for r, v in nonempty_rows(ws, 4, ws.max_row, 9):
        cert, val, date, adv, ex, lead, school, period, score = v
        filled = [x for x in (cert, adv, school) if x]
        if len(filled) > 1:
            rep.error("其他", f"行{r} 一行只能填一类(证书/宿舍/海外)")
        if cert:
            if cert not in CERT_NAMES:
                rep.error("其他", f"行{r} 证书[{cert}]非法")
            if not val:
                rep.error("其他", f"行{r} 证书必须填分数/等级")
        if adv is not None and adv not in ("是", "否"):
            rep.error("其他", f"行{r} 先进宿舍应为 是/否")
        if lead == "是" and adv != "是" and ex != "是":
            rep.warn("其他", f"行{r} 宿舍长+1需宿舍先获称号")

    print(f"===== 校验报告: {Path(args.file).name} =====")
    print(f"ERROR {len(rep.errors)} 项 / WARN {len(rep.warns)} 项")
    for e in rep.errors:
        print(f"  ✗ {e}")
    for w in rep.warns:
        print(f"  △ {w}")
    if not rep.errors and not rep.warns:
        print("  全部通过。")
    sys.exit(2 if rep.errors else 0)


if __name__ == "__main__":
    main()
