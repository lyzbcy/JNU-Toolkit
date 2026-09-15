# -*- coding: utf-8 -*-
"""
综测个人加分项信息统计表 自动填写脚本
用法:
  python fill_form.py --data 材料.json --template 模板.xlsx --dry-run        # 试算，不写表
  python fill_form.py --data 材料.json --template 模板.xlsx --out 产出.xlsx   # 写入
算分规则依据《人工智能与计算机学院本科生综合素质测评实施细则》(2025年7月修订)。
本脚本是算分的唯一权威实现; 不确定的条目打 [待确认] 标记, 以班长/辅导员核定为准。
"""
import argparse
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from copy import copy as _style_copy

from openpyxl import load_workbook


def _put(w, r, c, v):
    """写入单元格并重置字体颜色为默认黑。

    模板的示例行(例/XXX)是红字, 直接写值会继承红色示例格式;
    真实数据必须以正常黑字呈现, 故写值时统一剥离显式字体颜色。
    """
    cell = w.cell(row=r, column=c)
    cell.value = v
    f = cell.font
    if f is not None and f.color is not None:
        nf = _style_copy(f)
        nf.color = None
        cell.font = nf
    return cell

# ---------------------------------------------------------------- 分值规则

COMP_TIERS = {"特等奖": 0, "一等奖": 1, "二等奖": 2, "三等奖": 3}
COMP_TABLE = {  # 级别: [特等, 一等, 二等, 三等]; None = 细则未明确, 需班长核定
    "A1": [15, 13, 11, 9],
    "A2": [11, 9, 8, None],
    "A3": [8, 6, 5, None],
    "B1": [5, 3, 2, None],
    "B2": [2, 1, 0, 0],
}
COMP_CATEGORIES = [
    "算法与程序设计", "数据挖掘与人工智能", "信息可视化", "物联网与嵌入式系统",
    "数学建模", "网络安全", "计算机体系结构", "项目应用与创新创业类",
]
KECHUANG_TYPES = ["已参与两次讲座类", "已参与一次以上学科竞赛"]
AWARD_LEVELS = ["国家级", "省级", "市级", "校级"]

PAPER_TABLE = {
    "SCI/SCIE一区期刊论文": 11, "SCI/SCIE二区期刊论文": 11,
    "SCI/SCIE三区期刊论文": 9, "SCI/SCIE四区期刊论文": 9,   # ⚠️ 细则未分区, 估档
    "EI期刊论文": 9, "SCI/SCIE收录的会议论文": 9,
    "CCF-A类会议论文": 9,                                  # ⚠️ 细则只写CCF会议=6, A类上浮估
    "CCF-B类会议论文": 6, "CSCD核心库期刊论文": 6,
    "国家公布的核心期刊": 4, "国际学术会议论文": 4,
    "国内普通期刊": 2, "国内学术会议发表的论文": 2, "国家有关出版社公开出版的著作": 2,
    "授权发明专利": 6, "授理发明专利": 2,
}
PAPER_FLAGGED = {"SCI/SCIE三区期刊论文", "SCI/SCIE四区期刊论文", "CCF-A类会议论文"}

DACHUANG_TABLE = {
    "国家级大创创业项目": 9, "国家级大创创新项目": 6,
    "省级大创创业项目": 6, "省级大创创新项目": 5,
    "校级大创": 2, "院级大创": 1,
}

PRACTICE_TABLE = {
    "省级以上重点团队": {"队员": 4, "队长": 5},
    "校级重点团队": {"队员": 3, "队长": 4},
    "院级重点团队": {"队员": 2, "队长": 3},
    "院级一般团队": {"队员": 1, "队长": 2},
}

WENYI_TABLE = {"省部级以上": [5, 4, 3], "市级": [4, 3, 2], "校级": [3, 2, 1], "院级": [2, 1, 0]}
TIYU_GRADE_IDX = {"第1至第3名": 0, "第4至第6名": 1, "第7至8名": 2}
HONOR_TABLE = {"国家级": 4, "省级": 3, "市级": 2, "校级": 1}
ACT_TYPES = ["文艺类", "体育类", "志愿类", "其他荣誉"]
ACT_LEVELS = ["省部级以上", "市级", "校级", "院级"]

CERT_NAMES = ["CET-4", "CET-6", "雅思", "托福", "CSP", "计算机软考"]


def cert_score(name, val):
    """返回 (分数, 待确认说明或None)。边界值就高计。"""
    try:
        s = float(str(val).strip())
    except ValueError:
        s = None
    if name == "计算机软考":
        v = str(val).strip()
        if "高" in v:
            return 4, None
        if "中" in v:
            return 3, None
        return None, f"软考等级[{val}]无法识别(应为 中级/高级)"
    if s is None:
        return None, f"{name} 分数[{val}]不是数字"
    if name == "CET-4":
        if s < 425: return None, "CET-4 低于425不加分"
        return (1 if s < 500 else 2 if s < 560 else 3), None
    if name == "CET-6":
        if s < 425: return None, "CET-6 低于425不加分"
        return (2 if s < 500 else 3 if s < 560 else 4), None
    if name == "雅思":
        if s < 5: return None, "雅思低于5不加分"
        return (2 if s < 6 else 3 if s < 7 else 4), None
    if name == "托福":
        if s < 40: return None, "托福低于40不加分"
        return (2 if s < 65 else 3 if s < 90 else 4), None
    if name == "CSP":
        if s < 200: return None, "CSP低于200不加分"
        return (2 if s < 300 else 3 if s < 400 else 4), None
    return None, f"未知证书 {name}"


def share_coeff(rank):
    """合作分摊系数 1/(2n); 独立/未填 = 1。"""
    if rank in (None, "", 0, 1):
        return 1.0
    return 1.0 / (2 * int(rank))


def norm_text(s):
    """归一化: 去空白、全角转半角、统一引号, 用于竞赛名模糊匹配。"""
    s = unicodedata.normalize("NFKC", str(s))
    s = re.sub(r"\s+", "", s)
    return s.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")


# ---------------------------------------------------------------- 数据校验与算分

class Calculator:
    def __init__(self, comp_catalog):
        self.comp_catalog = comp_catalog          # {精确名: 级别}
        self.norm_catalog = {norm_text(k): (k, v) for k, v in comp_catalog.items()}
        self.rows = []    # 输出行: (sheet, row, values dict)
        self.items = []   # 汇总条目: dict(kind, desc, base, coeff, score, flag)
        self.errors = []

    def err(self, msg):
        self.errors.append(msg)

    def add_item(self, kind, desc, base, coeff=1.0, score=None, flag=None):
        if score is None and base is not None:
            score = round(base * coeff, 2)
        self.items.append(dict(kind=kind, desc=desc, base=base,
                               coeff=coeff, score=score, flag=flag))

    def match_competition(self, name):
        if name in self.comp_catalog:
            return name, self.comp_catalog[name]
        n = norm_text(name)
        if n in self.norm_catalog:
            return self.norm_catalog[n]
        for nk, (orig, lv) in self.norm_catalog.items():
            if nk in n or n in nk:
                return orig, lv
        return None, None

    # ---- 各板块
    def calc_basic(self, d):
        b = d.get("basic") or {}
        hours = b.get("second_classroom_hours")
        if hours is None:
            self.err("basic.second_classroom_hours 缺失(没有学时也请填0)")
        else:
            hours = min(int(hours), 40)
        kt = b.get("kechuang_type")
        if kt not in KECHUANG_TYPES:
            self.err(f"basic.kechuang_type [{kt}] 应为二选一: {KECHUANG_TYPES}")
        comps = b.get("competitions_participated") or []
        if kt == "已参与一次以上学科竞赛" and not comps:
            self.err("选了'已参与一次以上学科竞赛'但没提供竞赛名")
        self.basic_rows = (hours, kt, comps)
        self.add_item("基础分", f"第二课堂学时 {hours}(≥40学时得满分30, 由班长汇总计分)", None)

    def calc_competitions(self, d):
        for i, c in enumerate(d.get("competitions") or []):
            name = c.get("name", "")
            exact, level = self.match_competition(name)
            if not exact:
                if c.get("not_in_catalog"):
                    self.add_item("科研", f"[竞赛] {name}: 不在学校竞赛目录, 留待班长认定",
                                  None, flag="不在竞赛目录, 级别与分数留空待认定")
                    self.rows.append(("加分项-竞赛", None, dict(name=name)))
                    continue
                self.err(f"competitions[{i}].name [{name}] 不在竞赛目录(Sheet1), 如确认真实存在请加 not_in_catalog=true")
                continue
            gl = c.get("award_level")
            gg = c.get("award_grade")
            if gl not in AWARD_LEVELS:
                self.err(f"competitions[{i}].award_level [{gl}] 应为 {AWARD_LEVELS}")
                continue
            if gg not in COMP_TIERS:
                self.err(f"competitions[{i}].award_grade [{gg}] 应为 特等奖/一等奖/二等奖/三等奖")
                continue
            cat = c.get("category")
            if cat not in COMP_CATEGORIES:
                self.err(f"competitions[{i}].category [{cat}] 应为8个分类之一(见form-spec)")
                continue
            base = COMP_TABLE[level][COMP_TIERS[gg]]
            coeff = share_coeff(c.get("team_rank"))
            flag = None
            if base is None:
                flag = f"{level} 三等奖档位细则未明确, 待班长核定"
            elif gg == "特等奖":
                flag = "特等奖按最高档计, 待班长核定"
            desc = f"[竞赛] {exact} {gl}{gg}({level})"
            self.add_item("科研", desc, base, coeff, flag=flag)
            self.rows.append(("加分项-竞赛", None, dict(
                name=exact, level=level, award_level=gl, award_grade=gg,
                team_rank=c.get("team_rank"), advisor=c.get("advisor", ""),
                desc=c.get("desc", ""), category=cat,
                score=None if base is None else round(base * coeff, 2))))

    def calc_papers(self, d):
        for i, p in enumerate(d.get("papers") or []):
            t = p.get("type")
            if t not in ("论文/著作", "专利"):
                self.err(f"papers[{i}].type 应为 论文/著作 或 专利")
                continue
            lv = p.get("level")
            if lv not in PAPER_TABLE:
                self.err(f"papers[{i}].level [{lv}] 不在可用级别选项中(模板冗余值 其他/受理/授权 勿用)")
                continue
            rank = p.get("rank")
            if t == "专利":
                if rank and int(rank) > 5:
                    self.add_item("科研", f"[专利] {p.get('title')}: 第{rank}位超出前5位不计分", 0)
                    self.rows.append(("加分项-论文著作专利", None, dict(
                        type=t, title=p.get("title"), venue=p.get("venue", ""), level=lv,
                        rank=rank, advisor=p.get("advisor", ""), score=0)))
                    continue
            else:
                limit = 5 if "SCI" in lv else 3
                if rank and int(rank) > limit:
                    self.add_item("科研", f"[论文] {p.get('title')}: 第{rank}位超出前{limit}位不计分", 0)
                    self.rows.append(("加分项-论文著作专利", None, dict(
                        type=t, title=p.get("title"), venue=p.get("venue", ""), level=lv,
                        rank=rank, advisor=p.get("advisor", ""), score=0)))
                    continue
            base = PAPER_TABLE[lv]
            coeff = share_coeff(rank)
            flag = "论文/分区认定以图书馆为准" if lv in PAPER_FLAGGED else "论文类均需班长核定收录情况"
            if p.get("second_unit"):
                coeff *= 0.3
                flag = (flag + "; 学校为第二完成单位×0.3").strip("; ")
            self.add_item("科研", f"[{t}] {p.get('title')}({lv}) 第{rank}位", base, coeff, flag=flag)
            self.rows.append(("加分项-论文著作专利", None, dict(
                type=t, title=p.get("title"), venue=p.get("venue", ""), level=lv,
                rank=rank, advisor=p.get("advisor", ""),
                score=round(base * coeff, 2))))

    def calc_practices(self, d):
        for i, s in enumerate(d.get("social_practices") or []):
            tc = s.get("team_class")
            role = s.get("role")
            if tc == "个人社会实践":
                role = "个人参与社会实践"
            if tc not in PRACTICE_TABLE and tc != "个人社会实践":
                self.err(f"social_practices[{i}].team_class [{tc}] 应为四类团队或 个人社会实践")
                continue
            if role not in ("队长", "队员", "个人参与社会实践"):
                self.err(f"social_practices[{i}].role [{role}] 应为 队长/队员/个人参与社会实践")
                continue
            base = 2 if tc == "个人社会实践" else PRACTICE_TABLE[tc][role]
            self.add_item("实践", f"[社会实践] {s.get('name')}({tc}/{role})", base)
            self.rows.append(("加分项-社会实践", "practice", dict(
                name=s.get("name"), team_class=(tc if tc != "个人社会实践" else None),
                role=role, desc=s.get("desc", ""), score=base)))

    def calc_dachuang(self, d):
        for i, x in enumerate(d.get("dachuang") or []):
            lv = x.get("level")
            if lv not in DACHUANG_TABLE:
                self.err(f"dachuang[{i}].level [{lv}] 应为六类大创之一")
                continue
            passed = x.get("passed", True)
            if not passed:
                self.add_item("科研", f"[大创] {x.get('title')}: 立项未结题不加分(全体-3/队长-5由班长处理)", None)
            else:
                self.add_item("科研", f"[大创] {x.get('title')}({lv}/{x.get('role')})",
                              DACHUANG_TABLE[lv])
            self.rows.append(("加分项-社会实践", "dachuang", dict(
                level=lv, role=x.get("role"), title=x.get("title"), passed=passed,
                score=DACHUANG_TABLE[lv] if passed else None)))

    def calc_affairs(self, d):
        for a in d.get("student_affairs") or []:
            sc = a.get("score", 3)
            self.add_item("实践", f"[学生事务] {a.get('position')}", sc,
                          flag="学生事务分由辅导员/评定人核定(0-5)")
            self.rows.append(("加分项-学生事务", None, dict(
                position=a.get("position"), period=a.get("period", ""),
                desc=a.get("desc", ""), score=sc)))

    def calc_activities(self, d):
        for i, a in enumerate(d.get("activities") or []):
            t = a.get("type")
            if t not in ACT_TYPES:
                self.err(f"activities[{i}].type [{t}] 应为 {ACT_TYPES}")
                continue
            lv = a.get("award_level")
            flag, base, extra_note = None, None, None
            if t == "文艺类":
                gg = {"特等奖": 0, "一等奖": 0, "二等奖": 1, "三等奖": 2}.get(a.get("award_grade"))
                if gg is None:
                    self.err(f"activities[{i}] 文艺类 award_grade 应为 特等/一/二/三等奖")
                    continue
                if lv not in WENYI_TABLE:
                    self.err(f"activities[{i}].award_level [{lv}] 应为 {ACT_LEVELS}")
                    continue
                base = WENYI_TABLE[lv][gg]
                if a.get("award_grade") == "特等奖":
                    flag = "特等奖按一等奖档计, 待班长核定"
            elif t == "体育类":
                gi = TIYU_GRADE_IDX.get(a.get("award_grade"))
                if gi is None or lv not in WENYI_TABLE:
                    self.err(f"activities[{i}] 体育类 award_grade 应为名次档, award_level 应为 {ACT_LEVELS}")
                    continue
                base = WENYI_TABLE[lv][gi]
            elif t == "志愿类":
                base = 1
                flag = "志愿分由学工办按规模核定(1-2)"
                lv = None
            else:  # 其他荣誉
                hl = a.get("honor_level")
                if hl not in HONOR_TABLE:
                    self.err(f"activities[{i}] 其他荣誉需 honor_level ∈ 国家级/省级/市级/校级")
                    continue
                base = HONOR_TABLE[hl]
                lv = "省部级以上" if hl in ("国家级", "省级") else hl
            score = base
            if a.get("represent_school") and t in ("文艺类", "体育类"):
                score = min(score + 2, 5)
                extra_note = "代表学校外出+2(封顶5)"
            self.add_item("实践", f"[{t}] {a.get('name')}" +
                          (f"({extra_note})" if extra_note else ""), score, flag=flag)
            vh = a.get("volunteer_hours")
            if t == "志愿类" and not vh and "献血" in str(a.get("name") or ""):
                vh = "—"
            self.rows.append(("加分项-文体志愿类", None, dict(
                type=t, name=a.get("name"), award_level=lv,
                award_grade=a.get("award_grade"), volunteer_hours=vh,
                time=a.get("time"), score=score)))

    def calc_others(self, d):
        o = d.get("others") or {}
        for c in o.get("certificates") or d.get("certificates") or []:
            name = c.get("name")
            if name not in CERT_NAMES:
                self.err(f"certificates {c}: name 应为 {CERT_NAMES}")
                continue
            sc, fl = cert_score(name, c.get("score_or_grade"))
            self.add_item("其他", f"[证书] {name}={c.get('score_or_grade')}", sc, flag=fl)
            self.rows.append(("加分项-其他", "cert", dict(
                name=name, val=c.get("score_or_grade"), date=c.get("date", ""), score=sc)))
        dorm = o.get("dorm") or d.get("dorm") or {}
        adv, ex, lead = dorm.get("advanced"), dorm.get("exempt"), dorm.get("leader")
        ds = 0
        if adv or ex:
            ds = 2 + (1 if lead else 0)
            self.add_item("其他", f"[宿舍] 先进/免检宿舍{'+宿舍长' if lead else ''}", ds,
                          flag="宿舍分0-2由学院核定, 宿舍长+1需宿舍先获称号")
        elif lead:
            self.add_item("其他", "[宿舍] 宿舍长(宿舍未获称号, 不加分)", 0)
        if adv or ex or lead:
            self.rows.append(("加分项-其他", "dorm", dict(
                advanced=adv or False, exempt=ex or False, leader=lead or False, score=ds)))
        for v in o.get("overseas") or d.get("overseas") or []:
            kind = v.get("kind", "游学")
            sc = 2 if kind == "交换" else 1
            self.add_item("其他", f"[海外{kind}] {v.get('school')}", sc,
                          flag=None if kind == "交换" else "游学每次1分累计≤2")
            self.rows.append(("加分项-其他", "overseas", dict(
                school=v.get("school"), period=v.get("period"), score=sc)))


CAPS = {"科研": (20, 15), "实践": (15, 8), "其他": (7, 7)}


def summarize(items):
    totals = {}
    for it in items:
        if it["score"] is not None:
            totals[it["kind"]] = totals.get(it["kind"], 0) + it["score"]
    lines = []
    for k, (cap_rule, cap_form) in CAPS.items():
        t = round(totals.get(k, 0), 2)
        lines.append(f"  {k}: {t} 分  (细则上限{cap_rule} / 班级汇总表口径{cap_form})")
    return "\n".join(lines)


# ---------------------------------------------------------------- 写入

SHEET_START = {  # sheet: (首数据行, 清示例范围末行, 列数)
    "基础分": (4, 4, 6), "加分项-竞赛": (3, 3, 9), "加分项-论文著作专利": (3, 3, 7),
    "加分项-学生事务": (3, 3, 4), "加分项-文体志愿类": (4, 5, 7), "加分项-其他": (4, 5, 9),
}


def write_workbook(wb, calc, data, out_path):
    sid, name = data["student"]["sid"], data["student"]["name"]

    # 基础分
    ws = wb["基础分"]
    for c in range(1, 7):
        _put(ws, 4, c, None)
    _put(ws, 4, 1, str(sid))
    _put(ws, 4, 2, name)
    hours, kt, comps = calc.basic_rows
    _put(ws, 4, 3, hours)
    _put(ws, 4, 4, kt)
    for j, cp in enumerate(comps or []):
        exact, _ = calc.match_competition(cp.get("name", ""))
        _put(ws, 4 + j, 5, exact or cp.get("name", ""))
        _put(ws, 4 + j, 6, cp.get("time", ""))

    # 各加分 sheet: 先清示例行
    for sheet, (start, demo_end, ncols) in SHEET_START.items():
        if sheet == "基础分":
            continue
        w = wb[sheet]
        for r in range(start, demo_end + 1):
            for c in range(1, ncols + 1):
                _put(w, r, c, None)

    # 社会实践 sheet 双区结构单独清示例(行3-5)
    wsp = wb["加分项-社会实践"]
    for r in range(3, 6):
        for c in range(1, 6):
            _put(wsp, r, c, None)

    for sheet, zone, vals in calc.rows:
        if sheet == "加分项-竞赛":
            r = 3 + _count(wb, "加分项-竞赛", 3)
            w = wb["加分项-竞赛"]
            vals_list = [vals.get("name"), vals.get("level"), vals.get("award_level"),
                         vals.get("award_grade"), vals.get("team_rank"), vals.get("advisor"),
                         vals.get("desc", ""), vals.get("category"), vals.get("score")]
            if vals.get("level") is None:  # 不在目录: 只写名称与描述
                vals_list = [vals.get("name"), None, None, None, None, None,
                             vals.get("desc", ""), None, None]
            for c, v in enumerate(vals_list, 1):
                _put(w, r, c, v)
        elif sheet == "加分项-论文著作专利":
            r = 3 + _count(wb, "加分项-论文著作专利", 3)
            w = wb["加分项-论文著作专利"]
            for c, v in enumerate([vals["type"], vals["title"], vals["venue"], vals["level"],
                                   vals["rank"], vals["advisor"], vals["score"]], 1):
                _put(w, r, c, v)
        elif sheet == "加分项-社会实践":
            w = wb["加分项-社会实践"]
            if zone == "practice":
                r = 3 + _count(wb, "加分项-社会实践", 3, stop=31)
                row_vals = [vals["name"], vals["team_class"], vals["role"], vals["desc"], vals["score"]]
            else:  # 大创区从32行起, B列留空
                r = 32 + _count(wb, "加分项-社会实践", 32)
                status = "已结题通过" if vals["passed"] else "未结题"
                row_vals = [vals["level"], None, vals["role"], f"{vals['title']}({status})", vals["score"]]
            for c, v in enumerate(row_vals, 1):
                _put(w, r, c, v)
        elif sheet == "加分项-学生事务":
            r = 3 + _count(wb, "加分项-学生事务", 3)
            w = wb["加分项-学生事务"]
            for c, v in enumerate([vals["position"], vals["period"], vals["desc"], vals["score"]], 1):
                _put(w, r, c, v)
        elif sheet == "加分项-文体志愿类":
            r = 4 + _count(wb, "加分项-文体志愿类", 4)
            w = wb["加分项-文体志愿类"]
            for c, v in enumerate([vals["type"], vals["name"], vals["award_level"],
                                   vals["award_grade"], vals["volunteer_hours"],
                                   vals["time"], vals["score"]], 1):
                _put(w, r, c, v)
        elif sheet == "加分项-其他":
            w = wb["加分项-其他"]
            if "val" in vals:  # 证书行
                r = 4 + _count(wb, "加分项-其他", 4, key=lambda row: row[0].value in CERT_NAMES)
                for c, v in enumerate([vals["name"], vals["val"], vals["date"], None, None, None, None, None, vals["score"]], 1):
                    _put(w, r, c, v)
            elif "advanced" in vals:
                r = 4 + _count(wb, "加分项-其他", 4, key=lambda row: row[3].value in ("是", "否"))
                for c, v in enumerate([None, None, None, "是" if vals["advanced"] else "否",
                                       "是" if vals["exempt"] else "否", "是" if vals["leader"] else "否",
                                       None, None, vals["score"]], 1):
                    _put(w, r, c, v)
            else:
                r = 4 + _count(wb, "加分项-其他", 4, key=lambda row: row[6].value)
                for c, v in enumerate([None, None, None, None, None, None, vals["school"], vals["period"], vals["score"]], 1):
                    _put(w, r, c, v)

    wb.save(out_path)


def _count(wb, sheet, start, stop=None, key=None):
    """统计从 start 行起已写入的行数(用于顺序追加)。"""
    w = wb[sheet]
    n = 0
    r = start
    while r <= (stop or w.max_row):
        row = list(w[r])
        if key:
            if key(row):
                n += 1
        elif any(c.value is not None for c in row):
            n += 1
        r += 1
    return n


# ---------------------------------------------------------------- 主流程

def print_report(data, calc):
    sid, name = data["student"]["sid"], data["student"]["name"]
    print(f"\n===== 填报汇总: {sid} {name} =====")
    for it in calc.items:
        base = f"基准{it['base']}" if it["base"] is not None else "不计分"
        coeff = f" ×分摊{it['coeff']:.2f}" if it["coeff"] != 1.0 else ""
        sc = f"→ {it['score']}" if it["score"] is not None else "→ 待定"
        flag = f"  [待确认: {it['flag']}]" if it["flag"] else ""
        print(f"  {it['kind']:>4} | {it['desc']} | {base}{coeff} {sc}{flag}")
    print("\n----- 三项小计(个人表如实填, 超上限由班长汇总时截断) -----")
    print(summarize(calc.items))
    flags = [it for it in calc.items if it["flag"]]
    if flags:
        print(f"\n----- 待班长/辅导员核定 ({len(flags)}项) -----")
        for i, it in enumerate(flags, 1):
            print(f"  {i}. {it['desc']}: {it['flag']}")
    if calc.errors:
        print(f"\n----- 数据错误 ({len(calc.errors)}) -----")
        for e in calc.errors:
            print(f"  ! {e}")
        print("  请修正 材料.json 后重试。")
    print()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True, help="材料.json")
    ap.add_argument("--template", required=True, help="空白模板 xlsx")
    ap.add_argument("--out", help="产出 xlsx 路径")
    ap.add_argument("--dry-run", action="store_true", help="只试算不写表")
    args = ap.parse_args()

    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    if not data.get("student", {}).get("sid") or not data["student"].get("name"):
        print("错误: 材料.json 缺少 student.sid / student.name"); sys.exit(2)

    wb = load_workbook(args.template)
    ws1 = wb["Sheet1"]
    catalog = {}
    for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, max_col=2):
        if row[0].value and row[1].value:
            catalog[str(row[0].value)] = str(row[1].value)

    calc = Calculator(catalog)
    calc.calc_basic(data)
    calc.calc_competitions(data)
    calc.calc_papers(data)
    calc.calc_practices(data)
    calc.calc_dachuang(data)
    calc.calc_affairs(data)
    calc.calc_activities(data)
    calc.calc_others(data)

    print_report(data, calc)

    if calc.errors:
        sys.exit(2)
    if args.dry_run:
        print("（试算模式, 未写入表格）")
        return
    if not args.out:
        print("错误: 需要 --out 或 --dry-run"); sys.exit(2)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(wb, calc, data, out)
    print(f"已写入: {out}")
    flags_n = len([i for i in calc.items if i["flag"]])
    try:
        with open(out.parent / "综测填报日志.txt", "a", encoding="utf-8") as f:
            f.write(f"[{datetime.now():%Y-%m-%d %H:%M}] 写入 {out.name} "
                    f"sid={data['student']['sid']} 条目{len(calc.items)} 待核定{flags_n}\n")
    except Exception:
        pass
    print("请接着运行 validate_form.py 校验。")


if __name__ == "__main__":
    main()
